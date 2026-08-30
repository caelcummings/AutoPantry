#!/usr/bin/env python3
"""Export AutoPantry Recipe List.xlsx → AutoPantry/Resources/recipes.json."""

from __future__ import annotations

import json
import re
import ssl
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(
    "/Users/caelc/Library/CloudStorage/OneDrive-Personal/AutoPantry Recipe List.xlsx"
)
OUT = Path(__file__).resolve().parents[1] / "AutoPantry" / "Resources" / "recipes.json"

YES_VALUES = {"yes", "y", "true", "1", "snack"}


def head_ok(url: str, ctx: ssl.SSLContext, ua: dict) -> bool:
    try:
        req = urllib.request.Request(url, method="HEAD", headers=ua)
        with urllib.request.urlopen(req, context=ctx, timeout=12) as resp:
            return 200 <= resp.status < 400
    except Exception:
        try:
            req = urllib.request.Request(url, headers={**ua, "Range": "bytes=0-0"})
            with urllib.request.urlopen(req, context=ctx, timeout=12) as resp:
                return 200 <= resp.status < 400
        except Exception:
            return False


def pexels_photo_id(url: str):
    m = re.search(r"/photo/[^/]*?(\d+)/?$", url.strip())
    if m:
        return m.group(1)
    m = re.search(r"(\d{6,})/?$", url.strip())
    return m.group(1) if m else None


def resolve_image(raw, rid: str, existing_by_id: dict, ctx, ua):
    if rid in existing_by_id:
        return existing_by_id[rid]
    if not raw:
        return None
    raw = str(raw).strip()
    if "images.pexels.com" in raw:
        return raw
    pid = pexels_photo_id(raw) if "pexels.com" in raw else None
    if pid:
        candidates = [
            f"https://images.pexels.com/photos/{pid}/pexels-photo-{pid}.jpeg",
            f"https://images.pexels.com/photos/{pid}/pexels-photo-{pid}.jpg",
            f"https://images.pexels.com/photos/{pid}/pexels-photo-{pid}.png",
        ]
        for c in candidates:
            if head_ok(c, ctx, ua):
                return c
        return candidates[0]
    return raw if raw.startswith("http") else None


def parse_ingredients(text):
    if not text:
        return []
    parts = [p.strip() for p in str(text).split(",")]
    ingredients = []
    for part in parts:
        if not part:
            continue
        m = re.match(r"^(.*?)\s*\((.+)\)\s*$", part)
        if m:
            name, amount = m.group(1).strip(), m.group(2).strip()
        else:
            name, amount = part.strip(), ""
        if name:
            ingredients.append({"name": name, "amount": amount})
    return ingredients


def parse_instructions(text):
    if not text:
        return []
    text = str(text).strip()
    if not text:
        return []
    parts = re.split(r"\n+|\r+", text)
    steps = []
    for p in parts:
        p = re.sub(r"^\d+[\).]\s*", "", p.strip())
        if p:
            steps.append(p)
    if len(steps) <= 1 and ". " in text:
        steps = [s.strip() for s in re.split(r"(?<=\.)\s+", text) if s.strip()]
    return steps


def to_int(v):
    if v is None or v == "":
        return None
    try:
        return int(float(v))
    except Exception:
        return None


def to_float(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except Exception:
        return None


def parse_snack(v) -> bool:
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    text = str(v).strip().lower()
    return text in YES_VALUES


def col_index(header, *names):
    lowered = {}
    for i, h in enumerate(header):
        if h is None:
            continue
        key = str(h).strip().lower().rstrip("?")
        lowered[key] = i
    for name in names:
        normalized = name.lower().rstrip("?")
        if normalized in lowered:
            return lowered[normalized]
    return None


def main():
    existing_by_id = {}
    if OUT.exists():
        old = json.loads(OUT.read_text())
        for r in old.get("recipes", []):
            url = r.get("imageURL")
            if url and "images.pexels.com" in url:
                existing_by_id[r["id"]] = url

    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    header = [h for h in rows[1][1:]]
    print("headers", header)

    snack_idx = col_index(header, "snack", "isSnack", "is snack", "Snack")
    breakfast_idx = col_index(header, "breakfast", "isBreakfast", "is breakfast", "Breakfast?")
    if snack_idx is None:
        print("WARNING: no snack column found — all recipes will be isSnack=false")
    if breakfast_idx is None:
        print("WARNING: no breakfast column found — all recipes will be isBreakfast=false")

    ctx = ssl.create_default_context()
    ua = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
        )
    }

    def cell(cells, idx, default=None):
        if idx is None or idx >= len(cells):
            return default
        return cells[idx]

    recipes = []
    seen = set()
    for row in rows[2:]:
        cells = list(row[1 : 1 + len(header)])
        if not cells or not cells[0]:
            continue
        rid = str(cells[0]).strip()
        if rid in seen:
            continue
        seen.add(rid)

        title = str(cell(cells, 1) or "").strip() or rid.replace("-", " ")
        ingredients = parse_ingredients(cell(cells, col_index(header, "ingredients")))
        servings = to_int(cell(cells, col_index(header, "servingCount", "servings"))) or 1
        if servings < 1:
            servings = 1

        recipes.append(
            {
                "id": rid,
                "title": title,
                "imageURL": resolve_image(
                    cell(cells, col_index(header, "imageURL")),
                    rid,
                    existing_by_id,
                    ctx,
                    ua,
                ),
                "cookMinutes": to_int(cell(cells, col_index(header, "cookMinutes"))),
                "calories": to_int(cell(cells, col_index(header, "calories"))),
                "proteinGrams": to_int(cell(cells, col_index(header, "proteinGrams"))),
                "carbsGrams": to_int(cell(cells, col_index(header, "carbsGrams"))),
                "fatGrams": to_int(cell(cells, col_index(header, "fatGrams"))),
                "servings": servings,
                "ingredientCount": to_int(
                    cell(cells, col_index(header, "ingredientcount", "ingredientCount"))
                )
                or len(ingredients),
                "estimatedPrice": to_float(
                    cell(cells, col_index(header, "estimatedPrice"))
                ),
                "isSnack": parse_snack(cell(cells, snack_idx)),
                "isBreakfast": parse_snack(cell(cells, breakfast_idx)),
                "cuisines": [],
                "ingredients": ingredients,
                "instructions": parse_instructions(
                    cell(cells, col_index(header, "instructions"))
                ),
            }
        )

    payload = {
        "version": 3,
        "source": XLSX.name,
        "exportedRecipeCount": len(recipes),
        "recipes": recipes,
    }
    OUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n")
    snack_count = sum(1 for r in recipes if r["isSnack"])
    breakfast_count = sum(1 for r in recipes if r["isBreakfast"])
    print(f"wrote {OUT}")
    print(f"count {len(recipes)} · snacks {snack_count} · breakfasts {breakfast_count}")


if __name__ == "__main__":
    main()
